#!/usr/bin/env python3
"""Đọc sheet 03.Input_FS của model elevaTO (.xlsx) → sinh ai/js/targets/sheets.js.

03.Input_FS là "nơi duy nhất phải gõ số" của model, nên bố cục của nó CHÍNH LÀ bố cục
mà file xuất ra phải bám theo: nhãn dòng, thứ tự, số dòng, dòng nào là mục lớn / nhập tay /
công thức / dòng CHECK. Lấy thẳng từ template thay vì chép tay để file xuất ra dán được
vào model mà không lệch dòng nào.

    python3 tools/doc-template.py /đường/dẫn/elevaTO_Model_Template.xlsx > js/targets/sheets.js

Chạy lại mỗi khi template đổi; đừng sửa tay file sinh ra.
"""
import json
import re
import sys

import openpyxl

SHEET = '03.Input_FS'
# Bỏ lớp IF(C$4<=LastAct, …, "") của template: file xuất ra chỉ có năm actual nên luôn tính.
VO_LASTACT = re.compile(r'^=IF\([A-Z]+\$4<=LastAct,(.*),""\)$', re.S)
# Mục lớn: "A. ", "B1. ", "B2. " ở đầu nhãn. Mục con trong mục C dùng số La Mã ("I. ", "II. ").
MUC = re.compile(r'^([A-Z]\d?)\.\s')
LA_MA = re.compile(r'^(I{1,3}|IV|V)\.\s')


def ke_tiep(dang_o, moi):
    """Mục lớn chạy đúng thứ tự chữ cái, không nhảy bậc — nhờ đó "I. Cash flow…" (mục con
    La Mã trong mục C) không bị nhận lầm thành mục I."""
    if dang_o is None:
        return moi[0] == 'A'
    if moi[0] == dang_o[0]:
        return len(moi) > 1 and (len(dang_o) == 1 or int(moi[1]) == int(dang_o[1]) + 1)
    return ord(moi[0]) == ord(dang_o[0]) + 1


def dinh_dang(fmt):
    """Gộp number_format của template về vài lớp mà bộ ghi .xlsx tự dựng lại."""
    if '%' in fmt:
        return 'pct'
    if fmt.startswith('_('):
        return 'money'            # định dạng kế toán: âm trong ngoặc, 0 thành gạch ngang
    if '#,##0.0' in fmt:
        return 'dec'
    if '#,##0' in fmt or re.fullmatch(r'0+', fmt):
        return 'int'
    return 'text'


def phan_loai(lab, f, bold, co_so, la_head):
    if lab.startswith('→'):
        return 'check'
    if f:
        return 'calc'
    if not lab:
        return 'note'
    if la_head:
        return 'head'
    if bold and not co_so:
        return 'subhead'          # mục con: số La Mã trong mục C, tên nhóm TSCĐ trong mục F
    # Dòng có định dạng số ở cột C = dòng nhập tay; còn lại chỉ là chú thích.
    return 'input' if co_so else 'note'


def doc(path):
    wb = openpyxl.load_workbook(path)
    if SHEET not in wb.sheetnames:
        sys.exit(f'Template không có sheet {SHEET}; đang có: {", ".join(wb.sheetnames)}')
    ws = wb[SHEET]
    rows, muc = [], None
    for r in range(1, ws.max_row + 1):
        lab = ws.cell(r, 2).value
        lab = '' if lab is None else str(lab)
        raw = ws.cell(r, 3).value
        f = raw if isinstance(raw, str) and raw.startswith('=') else None
        if not lab and f is None:
            continue
        fmt = dinh_dang(ws.cell(r, 3).number_format or '')
        font = ws.cell(r, 2).font
        bold = bool(font and font.bold)
        m = MUC.match(lab)
        la_head = bool(m and bold and fmt == 'text' and ke_tiep(muc, m.group(1)))
        kind = phan_loai(lab, f, bold, fmt != 'text', la_head)
        if la_head:
            muc = m.group(1)
        # Nhãn ô B có thể là công thức trỏ sang nhãn khác (dòng 140–144 lấy tên mảng ở 132–136).
        ref = re.match(r'^=\$B\$(\d+)$', lab)
        rows.append({
            'r': r,
            'sec': muc,
            'label': '' if ref else lab,
            'ref': int(ref.group(1)) if ref else None,
            'kind': kind,
            # Công thức trong cùng sheet, đã bỏ lớp IF(…LastAct…) và dấu '=' đầu — dựng lại theo dòng mới.
            'f': cong_thuc(f) if f else None,
            'fmt': fmt,
        })
    kiem(rows)
    return rows


def cong_thuc(f):
    m = VO_LASTACT.match(f)
    return m.group(1) if m else f.lstrip('=')


def kiem(rows):
    """Dòng 1–5 do bộ dựng tự viết; từ dòng 6 trở đi công thức phải tự lập trong cùng sheet.

    Template mà thêm một công thức trỏ sang sheet khác hoặc còn dùng tên LastAct thì phải biết
    ngay lúc sinh file, chứ không để Excel từ chối mở file đã xuất ra.
    """
    xau = [r for r in rows if r['r'] > 5 and r['f'] and ('!' in r['f'] or 'LastAct' in r['f'])]
    if xau:
        sys.exit('Công thức ngoài phạm vi sheet ở dòng: '
                 + ', '.join(f"{r['r']} ({r['f'][:60]})" for r in xau))


def ra_js(rows):
    d = json.dumps(rows, ensure_ascii=False, separators=(',', ':'))
    return (
        '// SINH TỰ ĐỘNG bởi tools/doc-template.py từ elevaTO_Model_Template.xlsx — đừng sửa tay.\n'
        '// Bố cục sheet 03.Input_FS của model: r = số dòng trong model (giữ nguyên để dán được),\n'
        '// sec = mục lớn (A…I), kind = head | subhead | input | calc | check | note,\n'
        '// f = công thức trong cùng sheet (viết theo cột C), fmt = lớp định dạng số.\n'
        f'export const INPUT_FS = {d};\n'
    )


if __name__ == '__main__':
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    # Nhãn trong template có dấu tiếng Việt — đừng để locale của máy quyết định bảng mã.
    sys.stdout.reconfigure(encoding='utf-8')
    sys.stdout.write(ra_js(doc(sys.argv[1])))
